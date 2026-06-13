from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

navy = PatternFill('solid', fgColor='1B2A4A')
alt_fill = PatternFill('solid', fgColor='F2F6FC')
white_fill = PatternFill('solid', fgColor='FFFFFF')
hdr_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
data_font = Font(name='Arial', size=10)
bold_font = Font(name='Arial', size=10, bold=True)
wrap = Alignment(wrap_text=True, vertical='top')
thin_border = Border(
    left=Side(style='thin', color='D0D0D0'),
    right=Side(style='thin', color='D0D0D0'),
    top=Side(style='thin', color='D0D0D0'),
    bottom=Side(style='thin', color='D0D0D0')
)

def style_header(ws, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = navy
        cell.font = hdr_font
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        cell.border = thin_border
    ws.freeze_panes = 'A2'

def add_data_row(ws, row_idx, data, bold=False):
    fill = alt_fill if row_idx % 2 == 0 else white_fill
    f = bold_font if bold else data_font
    for c, val in enumerate(data, 1):
        cell = ws.cell(row=row_idx, column=c, value=val)
        cell.font = f
        cell.fill = fill
        cell.alignment = wrap
        cell.border = thin_border
    ws.row_dimensions[row_idx].height = 80

# ── SHEET 1: Test Cases ──
ws1 = wb.active
ws1.title = "Test Cases"
headers = ["Test Case ID","Test Suite","Use Case Ref","User Story Ref","Test Case Title","Description","Pre-conditions","Test Steps","Test Data","Expected Result","Post-conditions","Priority","Type","Status","Assigned To","Execution Date","Comments"]
for i, h in enumerate(headers, 1):
    ws1.cell(row=1, column=i, value=h)
widths = [15,18,15,15,30,40,35,50,35,40,30,10,12,10,15,15,25]
for i, w in enumerate(widths, 1):
    ws1.column_dimensions[get_column_letter(i)].width = w
style_header(ws1, len(headers))

tc = [
# ── Address Selection ──
["TC-CO-001","Address Selection","UC-4.3.1.1","US-401","Display saved addresses with labels","Verify checkout shows saved addresses with Home/Work/Other labels","User logged in; has saved addresses","1. Start checkout flow\n2. Observe address list\n3. Verify labels shown","User with 3 addresses","Addresses listed with labels (Home, Work, Other); default address pre-selected","Address selection screen displayed","High","Functional","Not Executed","","",""],
["TC-CO-002","Address Selection","UC-4.3.1.1","US-401","Default address pre-selected","Verify default address is pre-selected on checkout entry","User has a default address set","1. Enter checkout\n2. Check pre-selected address","Default address exists","Default address has radio button selected","Correct default selected","High","Functional","Not Executed","","",""],
["TC-CO-003","Address Selection","UC-4.3.2.1","US-401","No saved addresses prompt","Verify user prompted to add address when none exist","User has no saved addresses","1. Enter checkout\n2. Observe screen","No addresses","Prompt to add new address shown; cannot proceed without adding","User must add address","High","Functional","Not Executed","","",""],
["TC-CO-004","Address Selection","UC-4.3.1.1","US-401","Add new address during checkout","Verify user can add new address with form fields","User in checkout; no address or wants new one","1. Tap \"Add New Address\"\n2. Fill form (address line, city, state, pincode, phone, label)\n3. Submit","Valid address data","Address saved to DB encrypted; appears in address list","New address available","High","Functional","Not Executed","","",""],
["TC-CO-005","Address Selection","UC-4.3.2.2","US-401","Invalid pincode validation","Verify pincode must be 6 digits","User adding/editing address","1. Enter pincode \"123\" (3 digits)\n2. Submit form\n3. Enter pincode \"1234567\" (7 digits)","Invalid pincodes","Inline validation error; form cannot be submitted","Address not saved","High","Validation","Not Executed","","",""],
["TC-CO-006","Address Selection","UC-4.3.1.1","US-401","Select address and continue","Verify selected address stored in CheckoutContext","User on address screen; address selected","1. Select an address\n2. Tap \"Continue\"\n3. Verify navigation","Selected address","Address stored in CheckoutContext; navigated to Payment screen","CheckoutContext updated","High","Functional","Not Executed","","",""],
["TC-CO-007","Address Selection","UC-4.3.1.1","US-401","No address selected modal","Verify modal when no address selected and Continue tapped","User on address screen; no address selected","1. Deselect any address\n2. Tap Continue","No selection","Modal displayed prompting user to select an address","Cannot proceed","Medium","Validation","Not Executed","","",""],
["TC-CO-008","Address Selection","UC-4.4","US-401","Address encrypted at rest","Verify shipping addresses are encrypted in database","Admin DB access","1. Add address via checkout\n2. Query addresses table directly\n3. Check stored data","New address","Address data stored encrypted (encryptAddressPayload); decrypted on retrieval","Encryption enforced","High","Security","Not Executed","","",""],
["TC-CO-009","Address Selection","UC-4.3.1.1","US-401","Edit existing address","Verify user can edit a saved address","User has saved addresses","1. Tap edit on existing address\n2. Modify city\n3. Save","Updated city","Address updated; encrypted and saved; reflected in list","Address updated","Medium","Functional","Not Executed","","",""],
["TC-CO-010","Address Selection","UC-4.3.1.1","US-401","Delete address","Verify user can delete a saved address","User has multiple addresses","1. Tap delete on an address\n2. Confirm deletion","Address to delete","Address removed from list and DB","Address deleted","Medium","Functional","Not Executed","","",""],
# ── Payment Method ──
["TC-CO-011","Payment Method","UC-4.3.1.2","US-402","Display payment options","Verify all payment methods shown","User on Payment screen","1. Navigate to Payment screen\n2. Observe options","None","COD, ZPIN Wallet (with live balance), Razorpay (Card/UPI/Net Banking) displayed","Payment options visible","High","Functional","Not Executed","","",""],
["TC-CO-012","Payment Method","UC-4.3.1.2","US-402","Select COD payment","Verify COD selection allows proceeding","User on Payment screen","1. Select \"Cash on Delivery\"\n2. Tap Continue","paymentMethod=cod","No balance check needed; can proceed to summary","COD selected in context","High","Functional","Not Executed","","",""],
["TC-CO-013","Payment Method","UC-4.3.1.2","US-402","Select Wallet - sufficient balance","Verify Wallet selection with enough balance","User on Payment screen; wallet balance >= order total","1. Select \"ZPIN Wallet\"\n2. Verify balance shown\n3. Proceed","Wallet balance=2000, order=1500","Full amount payable via wallet; Continue button enabled","Wallet selected","High","Functional","Not Executed","","",""],
["TC-CO-014","Payment Method","UC-4.3.2.3","US-402","Select Wallet - insufficient balance","Verify Wallet warning when balance too low","User on Payment screen; wallet balance < order total","1. Select \"ZPIN Wallet\"\n2. Observe warning","Wallet balance=500, order=1500","Warning: balance, order total, shortfall (₹1000) shown; Continue disabled","Cannot proceed with wallet","High","Functional","Not Executed","","",""],
["TC-CO-015","Payment Method","UC-4.3.1.2","US-402","Select Razorpay","Verify Razorpay selection","User on Payment screen","1. Select Razorpay\n2. Proceed","paymentMethod=razorpay","Can proceed to summary; Razorpay SDK invoked on order placement","Razorpay selected","High","Functional","Not Executed","","",""],
["TC-CO-016","Payment Method","UC-4.3.1.2","US-402","Wallet live balance display","Verify wallet balance is fetched and shown live","User on Payment screen","1. Check ZPIN Wallet option\n2. Verify balance displayed","None","Live wallet balance shown next to Wallet option; decrypted from encrypted storage","Balance displayed","Medium","Functional","Not Executed","","",""],
# ── Coupon Application ──
["TC-CO-017","Coupon Application","UC-4.3.1.3","US-403","Apply valid percentage coupon","Verify percentage coupon calculates correct discount","User on summary screen; valid coupon exists","1. Enter coupon code\n2. Tap Apply\n3. Verify discount","code=SAVE20, discount_type=percentage, discount_value=20, orderTotal=1000","Discount = 20% of 1000 = ₹200; price breakdown updated; finalAmount reduced","Coupon applied","High","Functional","Not Executed","","",""],
["TC-CO-018","Coupon Application","UC-4.3.1.3","US-403","Percentage coupon with max cap","Verify max_discount caps percentage discount","Valid coupon: 20% off, max ₹150; orderTotal=1000","1. Apply coupon\n2. Check discount amount","20% of 1000=200, but max=150","Discount capped at ₹150 (not ₹200); Math.min(200, 150)=150","Cap enforced","High","Business Rule","Not Executed","","",""],
["TC-CO-019","Coupon Application","UC-4.3.1.3","US-403","Apply valid fixed amount coupon","Verify fixed discount applied","Valid fixed coupon: ₹100 off","1. Apply coupon code\n2. Verify discount","discount_type=fixed, discount_value=100","Discount = ₹100; capped at Math.min(100, totalAmount)","Fixed discount applied","High","Functional","Not Executed","","",""],
["TC-CO-020","Coupon Application","UC-4.3.2.4","US-403","Apply expired coupon","Verify expired coupon rejected","Coupon with valid_until in past","1. Enter expired coupon code\n2. Tap Apply","Expired coupon code","404: \"Invalid or expired coupon code\"","Coupon not applied","High","Validation","Not Executed","","",""],
["TC-CO-021","Coupon Application","UC-4.3.2.5","US-403","Coupon below minimum order amount","Verify min order check","Coupon min_order_amount=500; order total=300","1. Apply coupon\n2. Observe error","orderTotal=300, min=500","400: \"Minimum order value of ₹500 required for this coupon\"","Coupon not applied","High","Validation","Not Executed","","",""],
["TC-CO-022","Coupon Application","UC-4.3.2.6","US-403","Coupon usage limit reached","Verify used coupon rejected","Coupon usage_count >= usage_limit","1. Apply coupon that hit limit","Used coupon code","404: \"Invalid or expired coupon code\" (filtered by usage_count < usage_limit)","Coupon not applied","High","Validation","Not Executed","","",""],
["TC-CO-023","Coupon Application","UC-4.3.2.7","US-403","Invalid coupon code","Verify non-existent coupon rejected","No coupon with code \"FAKE123\"","1. Enter \"FAKE123\"\n2. Tap Apply","code=FAKE123","404: \"Invalid or expired coupon code\"","Coupon not applied","High","Validation","Not Executed","","",""],
["TC-CO-024","Coupon Application","UC-4.3.1.3","US-403","Remove applied coupon","Verify coupon can be removed","Coupon currently applied","1. Tap Remove/X on applied coupon\n2. Verify price updates","Applied coupon","Discount removed; finalAmount recalculated without coupon","Coupon removed","Medium","Functional","Not Executed","","",""],
["TC-CO-025","Coupon Application","UC-4.3.2.8","US-403","Coupon auto-invalidated on cart change","Verify coupon removed when total drops below minimum","Coupon applied; user removes item dropping total below min","1. Apply coupon (min_order=500, total=600)\n2. Remove item worth ₹200\n3. New total=400\n4. Verify coupon invalidated","Cart change causing total < min_order","Coupon automatically removed; discount cleared; user notified","Coupon auto-invalidated","High","Business Rule","Not Executed","","",""],
["TC-CO-026","Coupon Application","UC-4.4","US-403","Server-side coupon validation on order creation","Verify coupon re-validated server-side at order placement","Client sends couponCode; backend independently validates","1. POST /orders with couponCode\n2. Verify server ignores client couponDiscount\n3. Verify server queries coupons table","couponCode=SAVE20","Server validates coupon independently; never trusts client-supplied couponDiscount; discount calculated server-side","Security: no client manipulation","High","Security","Not Executed","","",""],
["TC-CO-027","Coupon Application","UC-4.3.1.3","US-403","Coupon validation rate limiter 20/15min","Verify couponLimiter on validation endpoint","User is logged in","1. Send 20 POST /coupons/validate requests in 15 min\n2. Send 21st","21 validation requests","First 20 succeed; 21st returns 429","Rate limit enforced","High","Security","Not Executed","","",""],
["TC-CO-028","Coupon Application","UC-4.3.1.3","US-403","Coupon usage count incremented on order","Verify usage_count +1 after successful order with coupon","Valid coupon with usage_count=5","1. Place order with coupon\n2. Check coupons table","couponCode with usage_count=5","After order: usage_count=6","Usage tracked","Medium","Functional","Not Executed","","",""],
# ── Browse Coupons ──
["TC-CO-029","Browse Coupons","UC-4.3.1.4","US-405","List all coupons with status","Verify GET /coupons returns all coupons with computed status","Coupons exist in DB","1. GET /api/v1/coupons\n2. Verify response","None","Returns coupons with id, code, description, discountType, discountValue, minOrderValue, maxDiscount, expiryDate, usageLimit, usageCount, status (active/expired/used)","Coupons listed","High","Functional","Not Executed","","",""],
["TC-CO-030","Browse Coupons","UC-4.3.2.11","US-405","Expired coupons shown with badge","Verify expired coupons displayed with Expired status","Expired coupons exist","1. GET /coupons\n2. Check expired ones","Coupon with valid_until in past","status='expired' for expired coupons; cannot be applied","Expired badge shown","Medium","Functional","Not Executed","","",""],
["TC-CO-031","Browse Coupons","UC-4.3.1.4","US-405","Used coupons shown with status","Verify used-up coupons shown","Coupon with usage_count >= usage_limit","1. GET /coupons\n2. Check used ones","Used coupon","status='used' when usage_count >= usage_limit","Used status shown","Medium","Functional","Not Executed","","",""],
["TC-CO-032","Browse Coupons","UC-4.3.1.4","US-405","Coupons cached in Redis","Verify coupons list cached with 1h TTL","Admin Redis access","1. GET /coupons\n2. Check Redis key coupons:all\n3. Verify TTL=3600","None","Cached with TTL_COUPONS=3600s","Cache reduces load","Medium","Performance","Not Executed","","",""],
# ── Order Summary ──
["TC-CO-033","Order Summary","UC-4.3.1.5","US-404","Order summary displays all sections","Verify summary screen shows complete order info","User completed address + payment steps","1. Navigate to Summary screen\n2. Verify all sections","Checkout context data","Shows: items with quantities/prices, shipping address, payment method, applied coupon, price breakdown (subtotal, shipping, GST 18%, coupon discount, final amount)","Complete summary shown","High","Functional","Not Executed","","",""],
["TC-CO-034","Order Summary","UC-4.3.1.5","US-404","Shipping fee calculation - free above ₹500","Verify free shipping when subtotal > 500","Cart subtotal = ₹800","1. Check shipping in price breakdown","subtotal=800","shippingAmount=0 (free shipping for orders > ₹500)","Free shipping applied","High","Business Rule","Not Executed","","",""],
["TC-CO-035","Order Summary","UC-4.3.1.5","US-404","Shipping fee calculation - ₹49 below ₹500","Verify ₹49 shipping when subtotal <= 500","Cart subtotal = ₹400","1. Check shipping in price breakdown","subtotal=400","shippingAmount=49","Shipping fee charged","High","Business Rule","Not Executed","","",""],
["TC-CO-036","Order Summary","UC-4.3.1.5","US-404","GST 18% calculation","Verify 18% GST calculated on subtotal","Cart subtotal = ₹1000","1. Check tax in price breakdown","subtotal=1000","taxAmount = Math.round(1000 * 0.18) = 180","GST calculated correctly","High","Business Rule","Not Executed","","",""],
["TC-CO-037","Order Summary","UC-4.3.1.5","US-404","Final amount calculation","Verify finalAmount = subtotal + shipping + tax - couponDiscount","subtotal=1000, shipping=0, tax=180, couponDiscount=100","1. Check finalAmount in breakdown","Known amounts","finalAmount = Math.max(0, 1000 + 0 + 180 - 100) = 1080","Correct final amount","High","Business Rule","Not Executed","","",""],
["TC-CO-038","Order Summary","UC-4.3.1.5","US-404","Delivery method toggle - instant trial","Verify instant trial adds ₹50 fee","User on summary screen","1. Toggle to Instant Trial delivery\n2. Check price breakdown","Instant trial selected","₹50 non-refundable trial fee added to breakdown; finalAmount increases by ₹50","Trial fee shown","High","Functional","Not Executed","","",""],
["TC-CO-039","Order Summary","UC-4.3.1.5","US-404","Expandable price breakdown","Verify price breakdown can be expanded","User on summary screen","1. Tap expandable breakdown\n2. Verify details","None","Subtotal, shipping, instant trial fee (if any), GST, coupon discount all shown","All line items visible","Medium","UI","Not Executed","","",""],
# ── Place Order ──
["TC-CO-040","Place Order","UC-4.3.1.5","US-701","Place order - successful COD","Verify order creation with COD payment","User completed checkout steps; payment=COD","1. Tap \"Place Order\"\n2. POST /api/v1/orders\n3. Verify response","paymentMethod=cod, shippingAddressId=valid","201: {success:true, order:{id, orderNumber:'ZP-xxx', finalAmount, status:'pending'}}; cart cleared","Order created, cart empty","High","Functional","Not Executed","","",""],
["TC-CO-041","Place Order","UC-4.3.1.5","US-701","Place order - required fields validation","Verify shippingAddressId and paymentMethod required","User logged in","1. POST /orders with missing shippingAddressId\n2. POST with missing paymentMethod","Missing required fields","400: {success:false, message:'shippingAddressId and paymentMethod are required'}","Order not created","High","Validation","Not Executed","","",""],
["TC-CO-042","Place Order","UC-4.3.1.5","US-701","Place order - empty cart rejected","Verify order creation fails with empty cart","User's cart is empty in DB","1. POST /orders with valid fields","Empty cart","400: {success:false, message:'Cart is empty'}; ROLLBACK","Order not created","High","Validation","Not Executed","","",""],
["TC-CO-043","Place Order","UC-4.3.1.5","US-701","Place order - invalid shipping address","Verify 404 for non-existent address","User logged in; addressId doesn't exist","1. POST /orders with bad shippingAddressId","Invalid shippingAddressId","404: {success:false, message:'Shipping address not found'}","Order not created","High","Validation","Not Executed","","",""],
["TC-CO-044","Place Order","UC-4.3.1.5","US-701","Place order - address ownership check","Verify address belongs to current user","User A's address ID used by User B","1. POST /orders with User A's addressId using User B's token","Cross-user addressId","404: Shipping address not found (WHERE user_id = $2)","Ownership enforced","High","Security","Not Executed","","",""],
["TC-CO-045","Place Order","UC-4.3.1.5","US-701","Place order - order number format","Verify order number is ZP-{orderId}","Successful order","1. Place order\n2. Check orderNumber format","None","orderNumber = 'ZP-' + 7-char id (e.g., 'ZP-A1B2C3D')","Format correct","Medium","Functional","Not Executed","","",""],
["TC-CO-046","Place Order","UC-4.3.1.5","US-701","Place order - inventory decremented","Verify product quantities decreased after order","Product with quantity=10; ordered qty=3","1. Place order with qty=3\n2. Check products table","Product quantity=10, order qty=3","Product quantity=7; if quantity reaches 0, in_stock set to false","Inventory updated","High","Functional","Not Executed","","",""],
["TC-CO-047","Place Order","UC-4.3.1.5","US-701","Place order - cart cleared after order","Verify cart emptied on successful order","User has items in cart","1. Place order\n2. GET /cart","Items in cart","Cart is empty after order; DELETE FROM cart_items WHERE user_id","Cart cleared","High","Functional","Not Executed","","",""],
["TC-CO-048","Place Order","UC-4.3.1.5","US-701","Place order - customer notification","Verify customer gets order confirmation notification","Successful order","1. Place order\n2. Check notifications table","None","Notification created: type='new_order', title='Order Confirmed: ZP-xxx', with order details in data","Notification sent","High","Functional","Not Executed","","",""],
["TC-CO-049","Place Order","UC-4.3.1.5","US-701","Place order - seller notifications","Verify each seller notified for their items","Order with items from 2 different sellers","1. Place order with multi-seller items\n2. Check notifications for each seller","Multi-seller order","Each unique seller gets notification with their items and total only","Seller notifications sent","High","Functional","Not Executed","","",""],
["TC-CO-050","Place Order","UC-4.3.1.5","US-701","Place order - rate limiter 10/hour","Verify orderCreateLimiter","User logged in","1. Place 10 orders in 1 hour\n2. Attempt 11th","11 order attempts","First 10 succeed; 11th returns 429 Too Many Requests","Rate limit enforced","High","Security","Not Executed","","",""],
["TC-CO-051","Place Order","UC-4.3.1.5","US-701","Place order - atomic transaction","Verify order creation uses DB transaction","Order creation process","1. Simulate failure mid-creation\n2. Verify ROLLBACK\n3. Check no partial data","Simulated failure","BEGIN/COMMIT pattern; on error ROLLBACK; no partial orders/items created","Atomicity enforced","High","Reliability","Not Executed","","",""],
["TC-CO-052","Place Order","UC-4.3.1.5","US-701","Place order - estimated delivery 5 days","Verify estimated_delivery set","Successful order","1. Place order\n2. Check estimated_delivery","None","estimated_delivery = NOW + 5 days (Asia/Kolkata timezone)","Delivery date set","Medium","Functional","Not Executed","","",""],
["TC-CO-053","Place Order","UC-4.3.1.5","US-701","Place order - client cart sync","Verify bodyCartItems synced when DB cart empty","DB cart empty; client sends cartItems array","1. POST /orders with cartItems in body\n2. Verify items synced to cart then order created","cartItems=[{productId,quantity}]","System syncs client items to DB cart, re-fetches, then creates order","Cart sync works","Medium","Functional","Not Executed","","",""],
# ── Wallet Payment ──
["TC-CO-054","Wallet Payment","UC-4.3.2.3","US-602","Wallet payment - sufficient balance","Verify wallet debit on order placement","Wallet balance=2000; order finalAmount=1500","1. POST /orders with paymentMethod=wallet\n2. Verify wallet debited","balance=2000, finalAmount=1500","Wallet balance reduced to 500; wallet_transaction created (type=debit); order status=confirmed, payment_status=paid","Wallet payment success","High","Functional","Not Executed","","",""],
["TC-CO-055","Wallet Payment","UC-4.3.2.10","US-602","Wallet payment - insufficient balance (race condition)","Verify error when balance drops between screens","Wallet balance=500; order finalAmount=1500","1. POST /orders with paymentMethod=wallet\n2. Check balance < finalAmount","balance=500 < 1500","400: {success:false, message:'Insufficient wallet balance'}; ROLLBACK","Order not created; balance unchanged","High","Security","Not Executed","","",""],
["TC-CO-056","Wallet Payment","UC-4.3.2.3","US-602","Wallet payment - wallet not found","Verify error when no wallet exists","User has no wallet row","1. POST /orders with paymentMethod=wallet","No wallet record","400: {success:false, message:'Wallet not found'}; ROLLBACK","Order not created","Medium","Validation","Not Executed","","",""],
["TC-CO-057","Wallet Payment","UC-4.3.2.3","US-602","Wallet payment - atomic within transaction","Verify wallet debit atomic with order creation","Wallet and order in same transaction","1. Place order with wallet\n2. Verify both in same BEGIN/COMMIT","Wallet payment","Wallet debit and order creation in same transaction; failure rolls back both","Atomicity enforced","High","Security","Not Executed","","",""],
["TC-CO-058","Wallet Payment","UC-4.3.2.3","US-602","Wallet payment - transaction record created","Verify wallet_transaction logged","Successful wallet payment","1. Place order with wallet\n2. Check wallet_transactions","Wallet payment","wallet_transaction: type='debit', amount=finalAmount, description='Payment for order ZP-xxx', reference_id=orderId","Transaction logged","Medium","Functional","Not Executed","","",""],
["TC-CO-059","Wallet Payment","UC-4.3.2.3","US-602","Wallet balance encrypted storage","Verify wallet balance stored/retrieved encrypted","Wallet operations","1. Check wallets table\n2. Verify balance is encrypted\n3. Verify decryptMoney works","None","Balance stored via encryptMoney(); retrieved via decryptMoney()","Encryption enforced","High","Security","Not Executed","","",""],
# ── Razorpay Payment ──
["TC-CO-060","Razorpay Payment","UC-4.3.1.2","US-603","Create Razorpay order","Verify POST /razorpay/create-order creates payment order","User has pending order","1. POST /razorpay/create-order with {orderId}\n2. Verify response","orderId of pending order","Returns: razorpayOrderId, amount (paise), currency='INR', keyId, appOrderId","Razorpay order created","High","Functional","Not Executed","","",""],
["TC-CO-061","Razorpay Payment","UC-4.3.1.2","US-603","Create Razorpay order - orderId required","Verify orderId is required","User logged in","1. POST /razorpay/create-order with {}","No orderId","400: {success:false, message:'orderId is required'}","Order not created","Medium","Validation","Not Executed","","",""],
["TC-CO-062","Razorpay Payment","UC-4.3.1.2","US-603","Create Razorpay order - order not found","Verify 404 for invalid orderId","User logged in","1. POST /razorpay/create-order with bad orderId","Non-existent orderId","404: {success:false, message:'Order not found'}","No Razorpay order","Medium","Validation","Not Executed","","",""],
["TC-CO-063","Razorpay Payment","UC-4.3.1.2","US-603","Create Razorpay order - already paid","Verify rejection for already-paid order","Order with payment_status=paid","1. POST /razorpay/create-order for paid order","Already paid order","400: {success:false, message:'Order is already paid'}","Duplicate prevented","Medium","Validation","Not Executed","","",""],
["TC-CO-064","Razorpay Payment","UC-4.3.1.2","US-603","Create Razorpay order - ownership check","Verify order belongs to current user","User A's order, User B's token","1. POST /razorpay/create-order with User A's orderId using User B's token","Cross-user orderId","404: Order not found (WHERE user_id = $2)","Ownership enforced","High","Security","Not Executed","","",""],
["TC-CO-065","Razorpay Payment","UC-4.3.1.2","US-603","Create Razorpay order - amount in paise","Verify amount converted to paise","Order finalAmount=1500.50","1. Create Razorpay order\n2. Check amount","finalAmount=1500.50","amount = Math.round(1500.50 * 100) = 150050 paise","Paise conversion correct","Medium","Functional","Not Executed","","",""],
["TC-CO-066","Razorpay Payment","UC-4.3.1.2","US-603","Verify Razorpay payment - valid signature","Verify HMAC-SHA256 signature verification","Payment completed in Razorpay","1. POST /razorpay/verify with valid razorpayOrderId, razorpayPaymentId, razorpaySignature, appOrderId","Valid signature","Order marked: payment_status='paid', status='confirmed', payment_id stored; payment_success notification created","Payment verified","High","Functional","Not Executed","","",""],
["TC-CO-067","Razorpay Payment","UC-4.3.1.2","US-603","Verify Razorpay - invalid signature","Verify invalid signature rejected","Tampered signature","1. POST /razorpay/verify with wrong signature","Invalid razorpaySignature","400: {success:false, message:'Invalid payment signature'}; HMAC mismatch","Payment not verified","High","Security","Not Executed","","",""],
["TC-CO-068","Razorpay Payment","UC-4.3.1.2","US-603","Verify Razorpay - required fields","Verify all 4 fields required for verification","User logged in","1. POST /razorpay/verify with missing fields","Missing razorpayOrderId/PaymentId/Signature/appOrderId","400: all four fields required message","Verification rejected","Medium","Validation","Not Executed","","",""],
["TC-CO-069","Razorpay Payment","UC-4.3.1.2","US-603","Verify Razorpay - replay attack prevention","Verify same payment_id cannot be used twice","Payment already verified","1. POST /razorpay/verify with already-used razorpayPaymentId","Duplicate payment_id","400: {success:false, message:'Payment already used'}; SELECT id FROM orders WHERE payment_id=$1","Replay prevented","High","Security","Not Executed","","",""],
["TC-CO-070","Razorpay Payment","UC-4.3.1.2","US-603","Verify Razorpay - already paid idempotent","Verify already-paid order returns success","Order already paid","1. POST /razorpay/verify for already-paid order","Already paid order","{success:true, message:'Payment already recorded'}; no double-update","Idempotent handling","Medium","Functional","Not Executed","","",""],
["TC-CO-071","Razorpay Payment","UC-4.3.1.2","US-603","Verify Razorpay - atomic transaction","Verify verification uses DB transaction","Payment verification","1. Verify payment\n2. Check BEGIN/COMMIT\n3. Simulate failure","None","Atomic: order update + notification in same transaction; ROLLBACK on failure","Atomicity enforced","High","Reliability","Not Executed","","",""],
["TC-CO-072","Razorpay Payment","UC-4.3.1.2","US-603","Razorpay - payment failure handling","Verify order stays pending on payment failure","Razorpay SDK dismissed or failed","1. Dismiss Razorpay payment sheet\n2. Verify order status","Payment dismissed","Order remains status='pending'; error modal shown; option to retry","Graceful failure handling","High","Functional","Not Executed","","",""],
["TC-CO-073","Razorpay Payment","UC-4.3.1.2","US-603","Razorpay - auto-cancel on failure","Verify order auto-cancelled on payment failure/dismissal","Payment failed in frontend","1. Payment fails\n2. Frontend auto-cancels order","Payment failure","Order cancelled automatically; user shown error with option to retry","Auto-cancel works","Medium","Functional","Not Executed","","",""],
# ── Order History ──
["TC-CO-074","Order History","UC-7.3.1","US-702","View order history paginated","Verify GET /orders returns paginated list","User has orders","1. GET /api/v1/orders?page=1&limit=20\n2. Verify pagination","page=1, limit=20","Orders returned sorted by created_at DESC; excludes pending+pending_payment orders; pagination object included","Order list displayed","High","Functional","Not Executed","","",""],
["TC-CO-075","Order History","UC-7.3.1","US-702","Order history - user isolation","Verify user only sees own orders","User A logged in","1. GET /orders as User A\n2. Verify only User A's orders","User A token","Only orders WHERE user_id = $1 returned","Isolation enforced","High","Security","Not Executed","","",""],
["TC-CO-076","Order History","UC-7.3.1","US-703","View order details","Verify GET /orders/:id returns full details","User has an order","1. GET /api/v1/orders/:id\n2. Verify fields","Valid orderId","Returns: id, orderNumber, status, paymentStatus, totalAmount, shippingAmount, taxAmount, finalAmount, shippingAddress (decrypted), paymentMethod, items[], createdAt","Full details shown","High","Functional","Not Executed","","",""],
["TC-CO-077","Order History","UC-7.3.1","US-703","Order details - not found","Verify 404 for non-existent order","User logged in","1. GET /orders/nonexistent","Invalid orderId","404: {success:false, message:'Order not found'}","Not found handled","Medium","Validation","Not Executed","","",""],
["TC-CO-078","Order History","UC-7.3.1","US-703","Order details - ownership check","Verify user can't view another's order","User A's orderId, User B's token","1. GET /orders/:id with wrong user","Cross-user orderId","404: Order not found (WHERE user_id = $2)","Ownership enforced","High","Security","Not Executed","","",""],
# ── Cancel Order ──
["TC-CO-079","Cancel Order","UC-7.3.2","US-704","Cancel pending order","Verify PATCH /orders/:id/cancel works for pending orders","Order in 'pending' status","1. PATCH /orders/:id/cancel with {reason:'Changed mind'}\n2. Verify response","reason='Changed mind'","Order status='cancelled'; cancellation record created; refundAmount returned","Order cancelled","High","Functional","Not Executed","","",""],
["TC-CO-080","Cancel Order","UC-7.3.2","US-704","Cancel confirmed order","Verify cancellation works for confirmed orders","Order in 'confirmed' status","1. PATCH /orders/:id/cancel with {reason:'Found cheaper'}","reason='Found cheaper'","Order status='cancelled'; status was IN ('pending','confirmed')","Order cancelled","High","Functional","Not Executed","","",""],
["TC-CO-081","Cancel Order","UC-7.3.2","US-704","Cancel order - reason required","Verify reason is required","Order cancellable","1. PATCH /orders/:id/cancel with {} (no reason)\n2. PATCH with reason=null","Missing reason","400: {success:false, message:'reason is required'}","Cancel rejected","High","Validation","Not Executed","","",""],
["TC-CO-082","Cancel Order","UC-7.3.2","US-704","Cancel order - non-cancellable status","Verify shipped/delivered orders can't be cancelled","Order in 'shipped' status","1. PATCH /orders/:id/cancel","Shipped order","400: {success:false, message:'Order cannot be cancelled'} (status NOT IN pending,confirmed)","Cancel rejected","High","Business Rule","Not Executed","","",""],
["TC-CO-083","Cancel Order","UC-7.3.2","US-704","Cancel order - row locking (FOR UPDATE)","Verify concurrent cancel prevention","Two concurrent cancel requests","1. Send 2 simultaneous PATCH /cancel\n2. Verify only 1 succeeds","Concurrent requests","SELECT ... FOR UPDATE prevents race; one succeeds, other gets 400","Race condition handled","High","Security","Not Executed","","",""],
["TC-CO-084","Cancel Order","UC-7.3.2","US-704","Cancel order - wallet refund","Verify wallet credit on paid order cancellation with wallet refund","Paid order; refundMethod=wallet","1. Cancel with refundMethod='wallet'\n2. Check wallet balance\n3. Check wallet_transactions","Paid order, refundMethod=wallet","Wallet credited with finalAmount; wallet_transaction: type='credit', description='Refund for cancelled order'","Refund credited","High","Functional","Not Executed","","",""],
["TC-CO-085","Cancel Order","UC-7.3.2","US-704","Cancel order - unpaid order no refund","Verify no refund for unpaid order","Order with payment_status='pending' (unpaid COD)","1. Cancel unpaid order\n2. Check refundAmount","Unpaid order","refundAmount=0; refundStatus='not_applicable'; payment_status unchanged","No refund needed","Medium","Business Rule","Not Executed","","",""],
["TC-CO-086","Cancel Order","UC-7.3.2","US-704","Cancel order - ownership check","Verify user can only cancel own orders","User A's order, User B's token","1. PATCH /orders/:id/cancel with User B's token","Cross-user order","400: Order cannot be cancelled (WHERE user_id = $2 fails)","Ownership enforced","High","Security","Not Executed","","",""],
["TC-CO-087","Cancel Order","UC-7.3.2","US-704","Cancel order - cancellation record","Verify order_cancellations table populated","Successful cancellation","1. Cancel order\n2. Check order_cancellations table","Valid cancellation","Record created with: order_id, user_id, reason, refund_method","Cancellation logged","Medium","Functional","Not Executed","","",""],
["TC-CO-088","Cancel Order","UC-7.3.2","US-704","Cancel order - atomic transaction","Verify cancel uses DB transaction","Cancellation process","1. Cancel order\n2. Verify BEGIN/COMMIT/ROLLBACK","None","All operations (status update, wallet credit, cancellation record) in single transaction","Atomicity enforced","High","Reliability","Not Executed","","",""],
# ── Return Order ──
["TC-CO-089","Return Order","UC-7.3.3","US-705","Submit return request - delivered order","Verify POST /orders/:id/return for delivered order","Order in 'delivered' status","1. POST /orders/:id/return with {reason:'Defective', orderItemIds:['item1']}\n2. Verify response","reason='Defective'","Return request created: status='requested'; returnId returned; order status='return_requested'","Return initiated","High","Functional","Not Executed","","",""],
["TC-CO-090","Return Order","UC-7.3.3","US-705","Return request - reason required","Verify reason is required","Delivered order","1. POST /orders/:id/return with {} (no reason)","Missing reason","400: {success:false, message:'reason is required'}","Return rejected","High","Validation","Not Executed","","",""],
["TC-CO-091","Return Order","UC-7.3.3","US-705","Return request - non-delivered order rejected","Verify return only for delivered orders","Order in 'confirmed' status","1. POST /orders/:id/return","Non-delivered order","400: {success:false, message:'Only delivered orders can be returned'}","Return rejected","High","Business Rule","Not Executed","","",""],
["TC-CO-092","Return Order","UC-7.3.3","US-705","Return request - duplicate active return","Verify existing active return returns current status","Return already in 'requested' or 'picked_up' status","1. POST /orders/:id/return again","Already active return","Returns existing return info: {returnId, status, message:'Return request is already active'}","No duplicate created","Medium","Business Rule","Not Executed","","",""],
["TC-CO-093","Return Order","UC-7.3.3","US-705","Return request - order not found","Verify 404 for non-existent order","User logged in","1. POST /orders/nonexistent/return","Invalid orderId","404: {success:false, message:'Order not found'}","Return not created","Medium","Validation","Not Executed","","",""],
["TC-CO-094","Return Order","UC-7.3.3","US-705","Return request - ownership check","Verify user can only return own orders","User A's order, User B logged in","1. POST /orders/:id/return with User B's token","Cross-user order","404: Order not found (WHERE user_id = $2)","Ownership enforced","High","Security","Not Executed","","",""],
# ── Return Status ──
["TC-CO-095","Return Status","UC-7.3.3","US-706","Get return status","Verify GET /orders/:id/return-status","Active return exists","1. GET /orders/:id/return-status\n2. Verify response","Order with return","Returns: returnId, orderId, orderNumber, returnStatus, reason, description, orderItemIds, requestedAt, statusHistory","Return status displayed","High","Functional","Not Executed","","",""],
["TC-CO-096","Return Status","UC-7.3.3","US-706","Return status - no return found","Verify 404 when no return request exists","Order without return request","1. GET /orders/:id/return-status","No return","404: {success:false, message:'No return request found for this order'}","Not found handled","Medium","Validation","Not Executed","","",""],
["TC-CO-097","Return Status","UC-7.3.3","US-706","Update return status - picked_up","Verify PATCH /orders/:id/return-status to picked_up","Return in 'requested' status; seller/admin role","1. PATCH with {status:'picked_up'}\n2. Verify response","nextStatus=picked_up","Status updated to picked_up; picked_up_at set; status_history appended","Status advanced","High","Functional","Not Executed","","",""],
["TC-CO-098","Return Status","UC-7.3.3","US-706","Update return status - refunded","Verify transition from picked_up to refunded","Return in 'picked_up' status; admin role","1. PATCH with {status:'refunded'}\n2. Verify response","nextStatus=refunded","Status=refunded; refunded_at set; order status='returned', payment_status='refunded'","Refund processed","High","Functional","Not Executed","","",""],
["TC-CO-099","Return Status","UC-7.3.3","US-706","Return status - invalid transition","Verify invalid status transitions rejected","Return in 'requested' status","1. PATCH with {status:'refunded'} (skipping picked_up)","Invalid transition","400: {success:false, message:'Invalid return status transition'}; only requested→picked_up and picked_up→refunded allowed","Transition enforced","High","Business Rule","Not Executed","","",""],
["TC-CO-100","Return Status","UC-7.3.3","US-706","Return status - invalid status value","Verify only picked_up and refunded accepted","Return exists","1. PATCH with {status:'cancelled'}\n2. PATCH with {status:'approved'}","Invalid status values","400: {success:false, message:'status must be picked_up or refunded'}","Invalid status rejected","Medium","Validation","Not Executed","","",""],
["TC-CO-101","Return Status","UC-7.3.3","US-706","Return status - customer forbidden","Verify customers cannot update return status","Customer role","1. PATCH /orders/:id/return-status as customer","Customer token","403: {success:false, message:'Forbidden'}","Customer blocked","High","Security","Not Executed","","",""],
["TC-CO-102","Return Status","UC-7.3.3","US-706","Return status - seller ownership check","Verify seller can only update returns for their products","Seller with no products in order","1. PATCH return-status as seller not owning order items","Seller without ownership","404: Order not found (seller ownership query fails)","Ownership enforced","High","Security","Not Executed","","",""],
["TC-CO-103","Return Status","UC-7.3.3","US-706","Return status history stored","Verify status_history JSONB updated","Return with status changes","1. Update status to picked_up\n2. Update to refunded\n3. Check status_history","Status transitions","status_history array contains [{status,at,by,role}] for each transition","Full history maintained","Medium","Functional","Not Executed","","",""],
# ── Order Tracking ──
["TC-CO-104","Order Tracking","UC-5.3.1.3","US-504","Track order - timeline","Verify GET /orders/:id/track returns timeline","Order in 'shipped' status","1. GET /orders/:id/track\n2. Verify trackingHistory","Shipped order","Returns orderId, orderNumber, trackingNumber, currentStatus, estimatedDelivery, trackingHistory with steps: pending→confirmed→processing→shipped","Timeline displayed","High","Functional","Not Executed","","",""],
["TC-CO-105","Order Tracking","UC-5.3.1.3","US-504","Track order - with return request","Verify tracking includes return status","Order with active return","1. GET /orders/:id/track\n2. Check return steps in timeline","Order with return","currentStatus='return_in_progress'; trackingHistory includes return_requested, picked_up (if applicable), refunded milestones","Return tracking shown","High","Functional","Not Executed","","",""],
["TC-CO-106","Order Tracking","UC-5.3.1.3","US-504","Track cancelled order","Verify cancelled order shows in timeline","Cancelled order","1. GET /orders/:id/track","Cancelled order","trackingHistory includes cancelled step with timestamp","Cancelled step shown","Medium","Functional","Not Executed","","",""],
["TC-CO-107","Order Tracking","UC-5.3.1.3","US-504","Track order - not found","Verify 404 for non-existent order","User logged in","1. GET /orders/nonexistent/track","Invalid orderId","404: {success:false, message:'Order not found'}","Not found handled","Medium","Validation","Not Executed","","",""],
["TC-CO-108","Order Tracking","UC-5.3.1.3","US-504","Track order - ownership check","Verify user can only track own orders","Cross-user orderId","1. GET /orders/:id/track with wrong user","User B's order","404: Order not found (WHERE user_id = $2)","Ownership enforced","High","Security","Not Executed","","",""],
# ── Order Confirm (Seller/Admin) ──
["TC-CO-109","Order Confirm","UC-7.3.1","US-503","Confirm order - admin/seller","Verify PATCH /orders/:id/confirm by admin/seller","Order in 'pending' status; seller role","1. PATCH /orders/:id/confirm\n2. Verify status changed","Pending order, seller token","Order status='confirmed'; notification sent to customer","Order confirmed","High","Functional","Not Executed","","",""],
["TC-CO-110","Order Confirm","UC-7.3.1","US-503","Confirm order - customer forbidden","Verify customers cannot confirm orders","Customer role","1. PATCH /orders/:id/confirm as customer","Customer token","403: {success:false, message:'Forbidden'}","Customer blocked","High","Security","Not Executed","","",""],
["TC-CO-111","Order Confirm","UC-7.3.1","US-503","Confirm order - seller ownership","Verify seller can only confirm orders with their products","Seller without products in order","1. PATCH /orders/:id/confirm as non-owning seller","Seller without matching products","404: Order not found or cannot be confirmed","Ownership enforced","High","Security","Not Executed","","",""],
["TC-CO-112","Order Confirm","UC-7.3.1","US-503","Confirm order - non-pending rejected","Verify only pending orders can be confirmed","Order in 'confirmed' status","1. PATCH /orders/:id/confirm on already-confirmed order","Already confirmed order","400: Order not found or cannot be confirmed (WHERE status='pending')","Non-pending rejected","Medium","Business Rule","Not Executed","","",""],
["TC-CO-113","Order Confirm","UC-7.3.1","US-503","Confirm order - customer notification","Verify customer notified on confirmation","Pending order confirmed","1. Confirm order\n2. Check notifications","Confirmed order","Notification: type='order_update', title='Order Confirmed: ZP-xxx', with orderId/status in data","Notification sent","Medium","Functional","Not Executed","","",""],
# ── Checkout Frontend ──
["TC-CO-114","Checkout Frontend","UC-4.3.1.5","US-404","Checkout flow - address to payment navigation","Verify multi-step navigation","User in checkout","1. Select address → Continue\n2. Verify Payment screen shown","Valid address selected","Navigation: address → payment selection screen; CheckoutContext persists","Multi-step works","High","Functional","Not Executed","","",""],
["TC-CO-115","Checkout Frontend","UC-4.3.1.5","US-404","Checkout flow - payment to summary navigation","Verify payment to summary step","User on Payment screen","1. Select payment method → Continue\n2. Verify Summary screen","Payment method selected","Navigation: payment → summary screen; all context preserved","Navigation works","High","Functional","Not Executed","","",""],
["TC-CO-116","Checkout Frontend","UC-4.3.1.5","US-404","Order confirmation screen","Verify success screen after order placed","Order successfully created","1. Place order\n2. Verify confirmation screen","Successful order","Order Confirmation screen shows order number and estimated delivery","Confirmation shown","High","UI","Not Executed","","",""],
["TC-CO-117","Checkout Frontend","UC-4.3.2.9","US-404","Out-of-stock at order placement","Verify stock error handled in frontend","Item goes out of stock between summary and Place Order","1. Place order\n2. Backend returns stock error\n3. Verify redirect","Out of stock item","Error shown; item flagged in cart; user redirected back to cart","Stock error handled","High","Functional","Not Executed","","",""],
["TC-CO-118","Checkout Frontend","UC-4.3.1.5","US-404","Checkout context persistence","Verify CheckoutContext preserves data across steps","User navigating checkout","1. Select address\n2. Go to payment\n3. Go back\n4. Verify address still selected","Address + payment selections","CheckoutContext maintains all selections across forward/back navigation","Context persisted","Medium","Functional","Not Executed","","",""],
]

for i, row in enumerate(tc):
    add_data_row(ws1, i + 2, row)

# ── SHEET 2: Execution Summary ──
ws2 = wb.create_sheet("Execution Summary")
es_headers = ["Test Suite","Total Test Cases","Passed","Failed","Blocked","Not Executed","Pass Rate (%)"]
for i, h in enumerate(es_headers, 1):
    ws2.cell(row=1, column=i, value=h)
es_widths = [25, 18, 10, 10, 10, 15, 15]
for i, w in enumerate(es_widths, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
style_header(ws2, len(es_headers))

es_data = [
    ["Address Selection", 10, 0, 0, 0, 10, "0%"],
    ["Payment Method", 6, 0, 0, 0, 6, "0%"],
    ["Coupon Application", 12, 0, 0, 0, 12, "0%"],
    ["Browse Coupons", 4, 0, 0, 0, 4, "0%"],
    ["Order Summary", 7, 0, 0, 0, 7, "0%"],
    ["Place Order", 14, 0, 0, 0, 14, "0%"],
    ["Wallet Payment", 6, 0, 0, 0, 6, "0%"],
    ["Razorpay Payment", 14, 0, 0, 0, 14, "0%"],
    ["Order History", 5, 0, 0, 0, 5, "0%"],
    ["Cancel Order", 10, 0, 0, 0, 10, "0%"],
    ["Return Order", 6, 0, 0, 0, 6, "0%"],
    ["Return Status", 9, 0, 0, 0, 9, "0%"],
    ["Order Tracking", 5, 0, 0, 0, 5, "0%"],
    ["Order Confirm", 5, 0, 0, 0, 5, "0%"],
    ["Checkout Frontend", 5, 0, 0, 0, 5, "0%"],
    ["TOTAL", 118, 0, 0, 0, 118, "0%"],
]
for i, row in enumerate(es_data):
    is_total = row[0] == "TOTAL"
    r = i + 2
    fill = alt_fill if r % 2 == 0 else white_fill
    f = bold_font if is_total else data_font
    for c, val in enumerate(row, 1):
        cell = ws2.cell(row=r, column=c, value=val)
        cell.font = f
        cell.fill = fill
        cell.alignment = wrap
        cell.border = thin_border
    ws2.row_dimensions[r].height = 30

# ── SHEET 3: Traceability Matrix ──
ws3 = wb.create_sheet("Traceability Matrix")
tm_headers = ["Requirement ID","Requirement Description","Test Case IDs","Coverage Status"]
for i, h in enumerate(tm_headers, 1):
    ws3.cell(row=1, column=i, value=h)
tm_widths = [18, 35, 70, 18]
for i, w in enumerate(tm_widths, 1):
    ws3.column_dimensions[get_column_letter(i)].width = w
style_header(ws3, len(tm_headers))

tm_data = [
    ["UC-4.3.1.1","Address Selection","TC-CO-001 to TC-CO-010","Covered"],
    ["UC-4.3.1.2","Payment Method Selection","TC-CO-011 to TC-CO-016, TC-CO-060 to TC-CO-073","Covered"],
    ["UC-4.3.1.3","Coupon Application","TC-CO-017 to TC-CO-028","Covered"],
    ["UC-4.3.1.4","Browse Available Coupons","TC-CO-029 to TC-CO-032","Covered"],
    ["UC-4.3.1.5","Order Summary & Place Order","TC-CO-033 to TC-CO-053","Covered"],
    ["UC-4.3.2.1","No Saved Addresses","TC-CO-003","Covered"],
    ["UC-4.3.2.2","Invalid Pincode","TC-CO-005","Covered"],
    ["UC-4.3.2.3","Insufficient Wallet Balance","TC-CO-014, TC-CO-055, TC-CO-056","Covered"],
    ["UC-4.3.2.4","Expired Coupon","TC-CO-020","Covered"],
    ["UC-4.3.2.5","Below Minimum Order","TC-CO-021","Covered"],
    ["UC-4.3.2.6","Usage Limit Reached","TC-CO-022","Covered"],
    ["UC-4.3.2.7","Invalid Coupon Code","TC-CO-023","Covered"],
    ["UC-4.3.2.8","Coupon Auto-Invalidated","TC-CO-025","Covered"],
    ["UC-4.3.2.9","Out-of-Stock at Placement","TC-CO-117","Covered"],
    ["UC-4.3.2.10","Wallet Race Condition","TC-CO-055, TC-CO-057","Covered"],
    ["UC-4.3.2.11","Expired Coupon on Browse","TC-CO-030","Covered"],
    ["UC-5.3.1.3","Track Order Delivery","TC-CO-104 to TC-CO-108","Covered"],
    ["UC-7.3.1","Order Lifecycle - Place/View","TC-CO-040 to TC-CO-053, TC-CO-074 to TC-CO-078","Covered"],
    ["UC-7.3.2","Order Cancellation","TC-CO-079 to TC-CO-088","Covered"],
    ["UC-7.3.3","Order Return","TC-CO-089 to TC-CO-103","Covered"],
    ["US-401","Address Selection","TC-CO-001 to TC-CO-010","Covered"],
    ["US-402","Payment Method Selection","TC-CO-011 to TC-CO-016","Covered"],
    ["US-403","Coupon Application","TC-CO-017 to TC-CO-028","Covered"],
    ["US-404","Order Summary & Place Order","TC-CO-033 to TC-CO-039, TC-CO-114 to TC-CO-118","Covered"],
    ["US-405","Browse Available Coupons","TC-CO-029 to TC-CO-032","Covered"],
    ["US-501","Select Standard Delivery","TC-CO-034, TC-CO-035","Covered"],
    ["US-502","Select Instant Trial Delivery","TC-CO-038","Covered"],
    ["US-503","Admin Confirms Order","TC-CO-109 to TC-CO-113","Covered"],
    ["US-504","Order Tracking","TC-CO-104 to TC-CO-108","Covered"],
    ["US-602","Pay with ZPIN Wallet","TC-CO-054 to TC-CO-059","Covered"],
    ["US-603","Pay with Razorpay","TC-CO-060 to TC-CO-073","Covered"],
    ["US-604","Wallet Credit on Cancellation","TC-CO-084, TC-CO-085","Covered"],
    ["US-701","Place an Order","TC-CO-040 to TC-CO-053","Covered"],
    ["US-702","View Order History","TC-CO-074, TC-CO-075","Covered"],
    ["US-703","View Order Details","TC-CO-076 to TC-CO-078","Covered"],
    ["US-704","Cancel an Order","TC-CO-079 to TC-CO-088","Covered"],
    ["US-705","Submit Return Request","TC-CO-089 to TC-CO-094","Covered"],
    ["US-706","Track Return Status","TC-CO-095 to TC-CO-103","Covered"],
]
for i, row in enumerate(tm_data):
    r = i + 2
    fill = alt_fill if r % 2 == 0 else white_fill
    for c, val in enumerate(row, 1):
        cell = ws3.cell(row=r, column=c, value=val)
        cell.font = data_font
        cell.fill = fill
        cell.alignment = wrap
        cell.border = thin_border
    ws3.row_dimensions[r].height = 40

# ── SHEET 4: Sign-Off ──
ws4 = wb.create_sheet("Sign-Off")
so_headers = ["Role","Name","Signature","Date","Comments"]
for i, h in enumerate(so_headers, 1):
    ws4.cell(row=1, column=i, value=h)
so_widths = [22, 25, 25, 18, 35]
for i, w in enumerate(so_widths, 1):
    ws4.column_dimensions[get_column_letter(i)].width = w
style_header(ws4, len(so_headers))

roles = ["Project Manager","QA Lead","Dev Lead","Business Analyst","Product Owner"]
for i, role in enumerate(roles):
    r = i + 2
    fill = alt_fill if r % 2 == 0 else white_fill
    for c in range(1, 6):
        cell = ws4.cell(row=r, column=c, value=role if c == 1 else "")
        cell.font = data_font
        cell.fill = fill
        cell.alignment = wrap
        cell.border = thin_border
    ws4.row_dimensions[r].height = 40

wb.save(r'C:\Zpin\ZPinEcom\CHECKOUT_ORDER_MANAGEMENT_TEST_CASES.xlsx')
print("DONE: CHECKOUT_ORDER_MANAGEMENT_TEST_CASES.xlsx created with 118 test cases")
